from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'Docs/Reports/Bao_cao_du_an_Pokemon_Boss_Arena.docx'
doc=Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
s=doc.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(2); s.bottom_margin=Cm(1.8); s.left_margin=Cm(2.2); s.right_margin=Cm(2.2)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',25),('Heading 1',18),('Heading 2',12.5)]:
    doc.styles[name].font.size=Pt(size); doc.styles[name].font.bold=True
    doc.styles[name].paragraph_format.space_before=Pt(10)
    doc.styles[name].paragraph_format.space_after=Pt(9)
doc.styles['Caption'].font.size=Pt(9); doc.styles['Caption'].font.italic=True
footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Pokémon Boss Arena  |  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
for r in footer.runs:r.font.size=Pt(8)
doc.core_properties.title='Báo cáo dự án Pokémon Boss Arena'
doc.core_properties.subject='Chức năng hiện tại cải tiến kỹ thuật kiểm thử và hướng phát triển'
doc.core_properties.author=''

def p(text,style=None):return doc.add_paragraph(text,style)
def h(text):doc.add_heading(text,2)
def page(title):doc.add_page_break();doc.add_heading(title,1)
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for c,v,w in zip(t.rows[0].cells,headers,widths):c.text=v;c.width=Cm(w)
    for vals in rows:
        for c,v,w in zip(t.add_row().cells,vals,widths):c.text=str(v);c.width=Cm(w)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if ri==0:repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,c in enumerate(row.cells):
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F3F5F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            pr.append(borders);m=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');m.append(e)
            pr.append(m)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1.05
                for run in para.runs:run.font.size=Pt(9);run.bold=ri==0
    p('')
def image(name,caption,width=15):
    para=p('');para.alignment=WD_ALIGN_PARAGRAPH.CENTER;para.paragraph_format.keep_with_next=True
    run=para.add_run();run.add_picture(str(ROOT/'Docs/Previews'/name),width=Cm(width))
    run._r.xpath('.//wp:docPr')[0].set('descr',caption)
    cap=p(caption,'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
def source(text):
    para=p('Căn cứ: '+text)
    for r in para.runs:r.font.size=Pt(8);r.italic=True

doc.add_heading('Báo cáo dự án Pokémon Boss Arena',0)
p('Chức năng hiện tại và quá trình cải tiến','Subtitle')
p('Ngày báo cáo 19 tháng 10 năm 2026')
h('Tóm tắt dự án')
p('Pokémon Boss Arena là trò chơi hành động 2D thời gian thực, chơi đơn ngoại tuyến trên Unity. Người chơi điều khiển Pikachu để đối đầu với Charizard, Blastoise và Venusaur trong ba đấu trường có chủ đề riêng. Trọng tâm của trải nghiệm là đọc tín hiệu tấn công, quản lý năng lượng, lựa chọn thời điểm dùng chiêu và né các đòn của boss.')
p('Bản hiện tại đã hình thành vòng chơi từ chọn boss, xuất hiện trong sân, chiến đấu đến thắng hoặc thua và chơi lại. Hệ thống có bốn kỹ năng của Pikachu, đánh thường, lướt, choáng, tương tác giữa các chiêu đối địch và chế độ luyện tập. Đây là bản thử nghiệm chức năng; âm thanh, cân bằng chuyên sâu, cơ chế năng lượng và choáng của boss vẫn cần hoàn thiện.')
image('Venusaur-PlayMode-HUD.png','Hình 1  Giao diện chơi thực tế ở màn Venusaur với tín hiệu Solar Beam',13.2)
h('Thông tin tổng quan')
table(['Thuộc tính','Giá trị hiện tại'],[
('Tên dự án','Pokémon Boss Arena'),('Công nghệ','Unity 6000.5.7f1 và C#'),('Đồ họa','Pixel art 2D với góc nhìn từ trên xuống'),('Nội dung','1 nhân vật điều khiển, 3 boss, 3 đấu trường'),('Trạng thái','Prototype có thể chơi và kiểm thử trong Unity')],[4,12.6])
source('README.md; Docs/REQUIREMENTS.md; mã nguồn trong Assets/Scripts; ảnh Play Mode của dự án.')

page('1 Mục tiêu phạm vi và yêu cầu')
h('Mục tiêu thiết kế')
p('Dự án xây dựng một hệ thống đấu boss trong đó chuyển động và thời điểm ra đòn có vai trò rõ ràng. Người chơi cần nhận biết nơi chiêu xuất phát, phạm vi nguy hiểm và khoảng thời gian còn lại để né. Mỗi boss tạo một bài toán chiến đấu khác nhau, thay vì chỉ thay đổi hình ảnh và lượng máu.')
h('Phạm vi đã triển khai')
table(['Nhóm','Chức năng'],[
('Vòng chơi','Chọn boss tự do; vào đấu trường; xuất hiện; chiến đấu; thắng hoặc thua; chơi lại; về menu.'),
('Người chơi','Pikachu, di chuyển tám hướng, camera bám theo, đánh thường, lướt và bốn kỹ năng.'),
('Chiến đấu','HP, năng lượng, hồi chiêu, tích sét, làm chậm, choáng và va chạm chiêu đối địch.'),
('Đối thủ','Charizard 500 HP; Blastoise 650 HP; Venusaur 750 HP, mỗi boss có bốn chiêu.'),
('Hỗ trợ thử nghiệm','Hình nhân; bật di chuyển hoặc miễn nhiễm; nạp sét; hồi hình nhân và khởi động lại trận.'),
('Giao diện','Thanh HP và năng lượng, thông tin boss, bốn ô chiêu, bộ đếm sét, vùng ngắm và cảnh báo trên mặt sân.')],[3.2,13.4])
h('Yêu cầu chất lượng')
p('Hoạt ảnh phải đúng hướng di chuyển và neo chân ổn định. Tín hiệu cảnh báo cần dễ đọc, môi trường không che nhân vật hoặc vùng sát thương. Kỹ năng chỉ tiêu hao tài nguyên sau khi đáp ứng điều kiện sử dụng; kết thúc trận và chơi lại phải xử lý sạch trạng thái chiến đấu.')
h('Ranh giới phiên bản')
p('Dự án tập trung vào các trận boss độc lập, chưa có quái thường, nhiệm vụ cốt truyện, nhiều người chơi hoặc hệ thống tài khoản. Ý tưởng bản đồ hành trình liên kết các màn chưa trở thành giao diện điều hướng hiện tại; người chơi đang sử dụng danh sách thẻ boss. Chưa có bằng chứng đo FPS, độ trễ đầu vào hoặc kiểm thử trên nhiều cấu hình máy để công bố chỉ tiêu hiệu năng.')
source('Docs/REQUIREMENTS.md; Docs/BOSS_SELECTION.md; Assets/Scripts/MainMenu.cs; Assets/Resources/BossCatalog.json.')

page('2 Điều khiển và bộ kỹ năng Pikachu')
p('Pikachu có 200 HP và tối đa 100 năng lượng. Khi vào sân, nhân vật chạy từ bên trái đến vòng trung tâm cùng hiệu ứng điện, sau đó mới mở điều khiển. Hoạt ảnh sử dụng tám hướng và điểm neo lấy từ dữ liệu bóng của sprite để hạn chế trượt hình.')
table(['Thao tác','Chức năng'],[
('W A S D','Di chuyển tám hướng; giữ hướng nhìn khi dừng.'),('Chuột trái','Bắn cầu điện nhỏ theo con trỏ; 5 sát thương, 5 năng lượng, hồi 1 giây; tầm 6 đơn vị.'),('Space','Lướt, tốn 10 năng lượng; không tạo miễn nhiễm sát thương.'),('Giữ E và nhấp chuột','Hiện vùng ngắm rồi chọn mục tiêu Thunderbolt.'),('Q và C','Q tăng tốc cho bản thân; C tự chọn đối thủ đủ gần.'),('Giữ X','Tích Electro Ball; có thể ném sớm bằng chuột trái.'),('Chuột giữa và Esc','Kéo camera tạm thời; Esc trở về menu.')],[4,12.6])
table(['Chiêu','Năng lượng','Hồi chiêu','Vận chiêu','Hiệu quả chính'],[
('E Thunderbolt','30','15 s','0,35 s','20 sát thương; tầm chọn 4; cộng 3 sét.'),
('Q Quick Attack','30','20 s','0,08 s','Hiệu lực 5 s; chạy và lướt ×1,5; vận chiêu ×1,3; cộng 2 sét.'),
('C Iron Tail','20','10 s','0,35 s','Tầm 1,8; 25 sát thương; chậm 40% trong 2 s; cộng 5 sét.'),
('X Electro Ball','70','60 s','Tối đa 1,5 s','20 đến 100 sát thương; cần và tiêu thụ 15 sét; tầm ngắm 8.')],[3,1.8,1.8,1.8,8.2])
h('Quy tắc chọn mục tiêu')
p('Thunderbolt khóa mục tiêu sống khi bắt đầu vận chiêu và bám mục tiêu đến lúc gây sát thương, nhưng có thể bị năng lượng đối địch cản hoặc giảm sức mạnh. Iron Tail tự chọn đối thủ gần, khóa điểm đánh và vẫn có thể hụt nếu đối thủ rời vùng đáp. Electro Ball khóa điểm đến khi ném, không tự đuổi theo mục tiêu.')
p('Năng lượng hồi 8 mỗi giây sau 2 giây kể từ lần tiêu hao, không hồi trong lúc vận chiêu. Sét được cộng khi hoàn thành kỹ năng thường, không tính theo số nhánh hiệu ứng hoặc số lần trúng; giới hạn là 15.')
source('Docs/PIKACHU_ABILITIES.md; Assets/Scripts/PikachuCombat.cs và PikachuPlayer.cs.')

page('3 Luật chiến đấu và phản hồi trạng thái')
h('Trúng đòn và choáng')
p('Khi nhận sát thương, Pikachu nháy đỏ trắng trong 0,28 giây và có tia va chạm. Nếu mất tổng cộng ít nhất 25 HP trong cửa sổ trượt 2 giây, nhân vật bị choáng 0,8 giây. Hoạt ảnh Hurt, ba ngôi sao quay quanh đầu và bộ đếm giúp người chơi phân biệt choáng với trúng đòn thông thường.')
p('Trong thời gian choáng, nhân vật không thể di chuyển, lướt, ngắm hoặc tấn công nhưng vẫn nhận sát thương. Chiêu đang vận bị hủy, hồi chiêu còn lại đặt thành 2 giây, không hoàn năng lượng hoặc sét đã tiêu. Sau choáng có 2 giây bảo vệ khỏi bị choáng tiếp; đây không phải miễn nhiễm sát thương.')
h('Hai chiêu đối địch va nhau')
p('Hệ thống so sánh sức mạnh còn lại của các đòn bay hoặc luồng năng lượng. Hai đòn ngang sức cùng mất; đòn mạnh hơn tiếp tục với sức mạnh bằng phần chênh lệch. Vụ nổ tại điểm chạm chỉ là hiệu ứng, không thêm sát thương diện rộng hoặc hố đất. Buff và đòn cận chiến không tham gia cơ chế này.')
table(['Tình huống','Kết quả'],[
('Điện thường 5 gặp đốm lửa 5','Cả hai bị triệt tiêu.'),('Electro Ball 100 gặp Air Slash 12','Cầu điện tiếp tục với 88 sát thương còn lại.'),('Electro Ball 100 gặp Fire Spin 48','Luồng lửa bị phá; cầu điện còn 52 sức mạnh.'),('Thunderbolt gặp năng lượng đối địch','Xử lý vật cản theo thứ tự trên đường phóng; có thể giảm hoặc chặn sát thương.')],[7,9.6])
h('Địa hình và khởi động lại')
p('Electro Ball tạo một lần sát thương vùng và một hố đất khi trúng hoặc chạm đất. Nếu bị chiêu khác triệt tiêu trước đó, nó không tạo hố. Hố là biến dạng thị giác có thể đi qua, tồn tại trong lượt chơi và chỉ được xóa khi nạp lại hoặc rời sân. Không có cơ chế rơi xuống hố.')
p('Chiến thắng hoặc thất bại dừng trận đấu. Chơi lại khôi phục HP, năng lượng và trạng thái chiến đấu theo scene mới. Chế độ hình nhân cho phép thử khóa mục tiêu, đòn hụt, làm chậm và miễn nhiễm mà không bị boss tấn công.')
source('Docs/SPELL_CLASHES.md; Docs/PIKACHU_ABILITIES.md; Assets/Scripts/SpellClash.cs.')

page('4 Boss Charizard và đấu trường lửa')
p('Charizard có 500 HP và đại diện cho nhóm thử thách né đạn, né lao tới và tránh luồng sát thương duy trì. Sân hiện tại là nền gạch xám hình chữ nhật, viền vàng và vòng giữa sân, có lửa trang trí. Đây chưa phải bản dựng núi lửa hoàn chỉnh; lửa trang trí không gây sát thương.')
table(['Chiêu','Sát thương','Hồi','Báo trước và cách hoạt động'],[
('Flamethrower','5 mỗi đốm','1 s','0,4 s; ba đốm lửa thẳng, cách nhau 0,12 s.'),
('Dragon Claw','25 một lần','8 s','0,8 s bay lên, vuốt phát sáng; lao tối đa 7 đơn vị trong 0,5 s.'),
('Air Slash','12 mỗi luồng','5 s','0,65 s; ba lưỡi khí ở góc −20, 0 và +20 độ.'),
('Fire Spin','8 mỗi 0,5 s','15 s','1,1 s; luồng vòng lửa mở rộng, duy trì 3 s với sáu nhịp.')],[3.2,2.4,1.4,9.6])
h('Logic và cách đối phó')
p('Các đòn khóa hướng từ khi bắt đầu cảnh báo. Người chơi di chuyển ra khỏi đường đạn, né ngang đường lao của Dragon Claw hoặc rời hình nón lửa trước khi Fire Spin phát huy hiệu lực. Fire Spin là luồng vòng lửa từ miệng boss theo hình tham chiếu, không phải vòng bẫy đặt quanh Pikachu.')
p('AI luân phiên các chiêu sẵn sàng, không vận hai chiêu đồng thời, nghỉ giữa các đòn và tiến gần khi khoảng cách lớn. Hiệu ứng làm chậm của Iron Tail ảnh hưởng tốc độ tiếp cận. Boss không tấn công trong màn xuất hiện của Pikachu.')
image('Charizard-Battle-HUD.png','Hình 2  Ảnh kiểm thử giao diện trận Charizard',14.5)
source('Docs/CHARIZARD_BOSS.md; Assets/Scripts/CharizardBoss.cs; ảnh Charizard-Battle-HUD.png.')

page('5 Boss Blastoise và đền nước ngập')
p('Blastoise có 650 HP. Sân đá tròn bao quanh bởi hồ xanh, rãnh nước nông, đền cổ và cột vỡ. Giới hạn di chuyển có bán kính 8,4 đơn vị. Mặt sân không trơn toàn bộ; trung tâm để thoáng cho việc quan sát tia nước và đường xoay.')
table(['Chiêu','Sát thương','Hồi','Vận chiêu','Cơ chế hiện tại'],[
('Hydro Pump','20 một lần','4 s','0,45 s','Hai tia từ pháo vai chụm thành luồng chính; bắn 2,4 s, xoay tối đa 40 độ/s.'),
('Water Pulse','Tối đa 24','6 s','0,35 s','Cầu nước đến điểm đã khóa; trúng trực tiếp 12, sóng lan 12.'),
('Rapid Spin','25 một lần','8 s','0,45 s','Xoay và tiến 4 đơn vị trong 0,65 s; có đường báo trước.'),
('Hydro Cannon','60 một lần','18 s','1,1 s','Nước từ miệng; bắn 0,8 s, bám mục tiêu tối đa 65 độ/s; luồng cong khi đổi hướng.')],[2.8,2,1.2,1.8,8.8])
h('Cải tiến hiệu ứng và va chạm')
p('Hydro Pump được sửa từ một luồng trước thân thành hai nhánh nối với pháo vai. Chỗ hợp tia được uốn thành chữ Y, vẽ các lớp màu đồng bộ để tránh đường viền cắt ngang lõi trắng. Bọt ở nòng, vệt dòng chảy và nước bắn tại đầu tia giúp phân biệt tia nước với tia điện.')
p('Hydro Cannon chuyển điểm phóng sang miệng, thu hẹp vùng báo còn 1,3 đơn vị chiều rộng. Bán kính luồng nước tăng từ 0,13 gần miệng đến 0,5. Hướng nước phía xa có độ trễ 0,32 giây so với hướng đầu, tạo đường cong khi boss đuổi theo. Hình vẽ và va chạm dùng chung đường cong; cảnh báo ban đầu chỉ biểu thị hướng bắt đầu, không bảo đảm vùng an toàn cố định trong suốt lúc bắn.')
h('Tương tác môi trường')
p('Bốn ô vòi nước được đánh dấu tại các trục của sân. Mỗi chu kỳ 10 giây kích hoạt một ô: báo trước 1,5 giây, phun 1 giây và gây tối đa 5 sát thương mỗi chu kỳ trong bán kính 0,85. Chúng tạm ngừng khi xuất hiện và sau khi trận kết thúc.')
source('Docs/BLASTOISE_ARENA.md; BlastoiseBoss.cs, BlastoiseVisuals.cs và WaterTempleEffects.cs. Thông số theo mã nguồn hiện tại.')

page('6 Boss Venusaur và khu vườn cổ')
p('Venusaur có 750 HP, được thêm vào danh sách boss thứ ba. Sân là vườn cổ bị rừng lấn, có lát đá, rêu, rễ cây, cột đổ và hoa lớn ở mép. Trung tâm rộng, phấn hoa nhẹ ở ngoài rìa; không có sương dày che tín hiệu tấn công.')
table(['Chiêu','Sát thương','Hồi','Báo trước và cách né'],[
('Razor Leaf','10 mỗi lá','3,5 s','0,55 s; năm lá cách nhau 20 độ, hướng khóa. Di chuyển qua khe lá.'),
('Leech Seed','6 mỗi nhịp','7 s','0,8 s; hạt bay vòng cung tới mảng đất bán kính 1,3. Rời vùng trước nhịp đầu.'),
('Vine Whip','25 một lần','5 s','0,65 s; quét cung 130 độ, tầm 3,2 trong 0,55 s. Lùi xa hoặc vòng ra sau.'),
('Solar Beam','60 một lần','16 s','1,65 s tích lực; tia khóa hướng, tầm 16, rộng 1,6; duy trì 0,8 s. Rời làn báo.')],[2.8,2.2,1.4,10.2])
p('Leech Seed gây năm nhịp cách nhau 0,6 giây, tổng tối đa 30 sát thương; nhịp đầu sau khi hạt đáp 0,3 giây, vùng hết hạn ở 3,2 giây. Razor Leaf và Solar Beam tham gia va chạm chiêu đối địch; dây leo và vùng đất không bị triệt tiêu theo quy tắc năng lượng.')
p('Hiệu ứng thể hiện nguồn chiêu: lá tập trung quanh tán, hạt đi từ hoa tới điểm đất, dây leo quét từ thân và Solar Beam tích sáng ở hoa. Các con số là cân bằng thử nghiệm cho trò chơi thời gian thực, không phải thông số chiến đấu theo lượt của nguyên tác.')
image('Venusaur-LeechSeed.png','Hình 3  Vùng Leech Seed và dây leo trong ảnh kiểm thử',13.4)
source('Docs/VENUSAUR_ARENA.md; Assets/Scripts/VenusaurBoss.cs và VenusaurVisuals.cs.')

page('7 Kiến trúc phần mềm và quản lý tài nguyên')
p('Dự án chia trách nhiệm giữa điều khiển, logic chiến đấu và hiển thị. Các boss kế thừa lớp ArenaBoss để dùng chung đối thủ, danh sách đạn, thông tin HP và giao diện va chạm luồng năng lượng. Mỗi boss vẫn tự xác định trình tự chiêu, hướng khóa, thời gian vận chiêu và điều kiện trúng đòn.')
table(['Thành phần','Trách nhiệm'],[
('MainMenu và BossCatalog.json','Đọc danh mục boss; hiển thị thẻ, phân trang, mở scene; mục chưa có scene hiện Sắp ra mắt.'),
('ArenaPreview và ArenaProgress','Sinh Pikachu và boss phù hợp, camera theo người chơi, chọn scene và giới hạn sân.'),
('PikachuPlayer và PikachuCombat','Chuyển động, hoạt ảnh, tài nguyên, kỹ năng, sát thương, trạng thái và điều kiện hành động.'),
('ArenaBoss và các lớp boss','Giao diện chung và logic riêng của Charizard, Blastoise, Venusaur.'),
('SpellClash','Mô phỏng đạn hai phía cùng dòng thời gian và giải quyết va chạm sớm nhất.'),
('CombatHud và các lớp Visuals','Giao diện, cảnh báo, trạng thái, hình đạn và hiệu ứng kỹ năng.'),
('PmdSpriteAnimation','Đọc metadata XML, thời lượng khung hình, tám hướng và điểm neo từ Shadow.'),
('CombatPixelSurface và môi trường','Vẽ hiệu ứng pixel và các lớp sân, phấn hoa hoặc gợn nước.'),
('Assets/Editor','Tạo scene, kiểm thử logic, chụp hình và chạy luồng Play Mode.')],[5,11.6])
h('Danh sách màn và tài nguyên')
p('Các scene chính gồm MainMenu, CharizardArena, BlastoiseArena và VenusaurArena. Danh mục boss lưu tên, đấu trường, scene, ảnh đại diện, màu và kích thước khung hình. Mỗi trang có bốn thẻ; thêm mục vào danh mục không tự tạo AI hoặc scene cho boss mới.')
p('Sprite được lấy từ PMDCollab, giữ metadata và credit tương ứng trong Resources. Cách nhập dùng lọc Point, không nén đối với sprite chiến đấu; dữ liệu Shadow cần đọc được để tính pivot. Việc giữ nguồn và credit hỗ trợ truy xuất tài nguyên, không thay thế việc rà soát quyền sử dụng trước khi phát hành.')
source('Assets/Scripts; Assets/Resources/BossCatalog.json; ProjectSettings/EditorBuildSettings.asset; Docs/PIKACHU_SOURCE.md.')

page('8 Quá trình cải tiến dự án')
p('Các cải tiến được thực hiện theo phản hồi trực tiếp trong quá trình chơi thử. Bảng dưới mô tả thay đổi về hành vi và giá trị sử dụng; không quy đổi thành tỷ lệ hiệu năng hoặc mức độ hài lòng khi chưa có đo lường.')
table(['Hạng mục','Vấn đề hoặc trạng thái trước','Cải tiến đã thực hiện'],[
('Hoạt ảnh Pikachu','Thứ tự hướng không khớp di chuyển; điểm neo thiếu ổn định.','Đối chiếu tám hàng sprite, neo chân từ Shadow và đồng bộ bước đi với quãng đường.'),
('Camera','Yêu cầu quan sát theo người chơi và không đổi góc khi dùng chiêu.','Camera bám vị trí mặt đất của Pikachu; không zoom khi ngắm; kéo chuột giữa để xem tạm.'),
('HUD kỹ năng','Bảng mô tả chiêu lớn chiếm diện tích sân.','Bốn icon nhỏ phía dưới, hồi chiêu dạng phủ; sét chỉ nằm dưới ulti.'),
('Ngắm và cận chiến','Người chơi khó nhận biết tầm dùng; C cần thao tác định hướng.','E có vòng giới hạn và tâm chọn mục tiêu; C tự chọn boss gần.'),
('Trúng đòn và choáng','Phản hồi khó nhận biết; cần ngưỡng sát thương theo thời gian.','Nháy sprite, tia va chạm, ngôi sao và khóa thao tác theo ngưỡng 25 HP trong 2 s.'),
('Va chạm kỹ năng','Đòn đối địch có thể xuyên nhau về mặt trải nghiệm.','So sánh sức mạnh, triệt tiêu hoặc giảm đòn; nổ chỉ tạo phản hồi hình ảnh.'),
('Blastoise','Luồng nước thiếu nguồn rõ ràng, nối tia gãy và ulti quá rộng.','Hai pháo hợp tia cong, ulti từ miệng, thu hẹp vùng và thêm bám mục tiêu có độ trễ.'),
('Chọn màn','Màn hai từng phụ thuộc chiến thắng; bố cục cố định hai boss.','Chọn tự do; danh mục dữ liệu và phân trang bốn thẻ để mở rộng.'),
('Lỗi khởi chạy menu','JsonUtility không có trong cấu hình module của dự án.','Bổ sung com.unity.modules.jsonserialize; Unity biên dịch và nạp lại code.'),
('Nội dung màn ba','Chưa có boss hệ cây và sân vườn.','Thêm Venusaur với bốn chiêu, vùng cảnh báo và kiểm thử scene riêng.')],[3,5.7,7.9])
source('Tài liệu trong Docs; thay đổi hiện tại trong mã nguồn; Packages/manifest.json; log Unity.')

page('9 Kiểm thử và bằng chứng xác nhận')
h('Phương pháp kiểm thử')
p('Dự án sử dụng kiểm tra xác định theo bước thời gian để xác minh điều kiện sử dụng, sát thương, hồi chiêu, phạm vi, đạn, trạng thái và va chạm. Play Mode được dùng để kiểm tra vòng đời Awake và Update, tải scene, xuất hiện, UI và luồng thắng thua. Ảnh dựng được xem để phát hiện sai hướng, lỗi vùng cảnh báo và lớp hiển thị.')
table(['Nhóm bằng chứng','Kết quả ghi nhận','Phạm vi và giới hạn'],[
('PikachuCombatCheck','100 kiểm tra trong log WaterTemplePlayMode','Kỹ năng, chi phí, buff, ulti, khóa mục tiêu và hố đất tại thời điểm chạy log.'),
('CharizardBossCheck','29 kiểm tra trong log hồi quy','Bốn chiêu, né, HP, AI và thắng thua.'),
('PikachuHitCheck','22 kiểm tra','Cửa sổ sát thương, ngưỡng choáng, bảo vệ, ngắt chiêu và dọn trạng thái.'),
('SpellClashCheck','20 kiểm tra','Đòn ngang hoặc lệch sức mạnh, đạn nhanh, thứ tự va chạm và quy tắc hố đất.'),
('Blastoise ban đầu','30 kiểm tra; tổng đợt hồi quy 201','Kết quả lịch sử trước một số thay đổi bám mục tiêu và rút ngắn vận chiêu.'),
('VenusaurCheck.log','23 kiểm tra thành công','Bốn chiêu, sát thương, khe né, mục tiêu khóa, vùng hết hạn, va chạm và giới hạn sân.'),
('VenusaurPlayMode.log','VENUSAUR_PLAYMODE_OK','Menu, vào scene, xuất hiện, tung chiêu, thắng và chơi lại trong Play Mode thực tế.')],[4.2,4.4,8])
h('Cách đọc kết quả')
p('Các con số trên là số điều kiện kiểm tra được khẳng định trong từng bộ thử, không phải số người chơi thử hoặc phần trăm bao phủ mã nguồn. Không cộng các lần chạy lặp thành tổng số kiểm thử độc lập. Đợt 201 kiểm tra là mốc lịch sử; các thay đổi Hydro Pump và Hydro Cannon mới hơn cần được chạy hồi quy đầy đủ trước khi chốt bản phát hành.')
p('Quá trình dựng Venusaur đã phát hiện lớp phấn hoa chưa khởi tạo làm màu sân nhạt đi; lớp này được xóa và áp dụng texture ngay khi tạo. Sau sửa, bộ 23 kiểm tra và luồng Play Mode đều hoàn tất. Chưa có benchmark đa cấu hình, báo cáo chơi dài hoặc khảo sát người dùng để đánh giá độ khó cuối cùng.')
source('Docs/Previews/WaterTemplePlayMode.log; WaterTempleFinalCheck.log; VenusaurCheck.log; VenusaurPlayMode.log.')

page('10 Hạn chế rủi ro và hướng phát triển')
h('Hạn chế hiện tại')
p('Bản thử chưa có nhạc và âm thanh chiến đấu. Tùy chọn âm lượng đã tồn tại nhưng không đồng nghĩa với việc nội dung âm thanh đã hoàn thành. Mana và choáng của boss vẫn là phần chưa triển khai; cơ chế choáng hiện được áp dụng cho Pikachu. Chưa có bản đồ hành trình liên kết, hệ thống lưu tiến trình toàn diện hoặc nhiều nhân vật điều khiển.')
p('Thông số sát thương và tốc độ xoay đang phục vụ thử nghiệm. Tia bám mục tiêu có thể tạo áp lực lớn khi kết hợp với choáng; cần đánh giá khả năng né bằng chơi thực tế. Môi trường và hiệu ứng tạo bằng mã có thể cần tối ưu khi số lượng đạn hoặc lớp pixel tăng; chưa có số đo CPU, GPU và bộ nhớ để kết luận về khả năng mở rộng.')
h('Kế hoạch ưu tiên đề xuất')
table(['Ưu tiên','Công việc','Tiêu chí hoàn thành đề xuất'],[
('P1','Chạy hồi quy toàn bộ sau thay đổi Blastoise.','Tất cả bộ kiểm tra chạy lại trên cùng bản mã; xác minh đường cong và thời gian cảnh báo mới.'),
('P1','Chơi thử cân bằng ba boss.','Ghi thời lượng trận, sát thương nhận, tỉ lệ né và tình huống khó tránh; điều chỉnh có căn cứ.'),
('P1','Đồng bộ tài liệu thông số.','Tài liệu và test đọc một nguồn dữ liệu cân bằng; loại bỏ mô tả lỗi thời về mở khóa.'),
('P2','Âm thanh và phản hồi kỹ năng.','Có âm riêng khi vận, phóng, trúng, va chiêu, choáng và thắng thua; âm lượng hoạt động.'),
('P2','Đo và tối ưu hiệu năng.','Có log profiler cho nhiều đạn và hiệu ứng; xem xét tái sử dụng đối tượng và giảm vẽ texture.'),
('P2','Mở rộng dữ liệu boss.','Tách cấu hình kỹ năng khỏi logic; dùng dữ liệu định nghĩa boss, giữ kiểm tra scene và ảnh.'),
('P3','Bổ sung nội dung và hành trình.','Boss mới có chủ đề, tín hiệu né, test; thiết kế lưu tiến trình và bản đồ nếu được chốt phạm vi.')],[1.6,5.3,9.7])
h('Điều kiện trước khi phát hành')
p('Cần thử bản build độc lập trên cấu hình mục tiêu, kiểm tra độ phân giải, dữ liệu lưu và khả năng phục hồi khi đổi scene. Nguồn sprite và hình tượng Pokémon cần được rà soát quyền sử dụng phù hợp với mục đích phát hành. Các mục này là việc cần làm tiếp, không phải kết quả đã được xác nhận.')
source('Phần hạn chế trong tài liệu kỹ năng và boss; đối chiếu mã nguồn hiện tại. Kế hoạch trên là đề xuất phát triển.')

page('11 Vận hành mở rộng và tài liệu tham chiếu')
h('Chạy và thử dự án')
p('Mở thư mục PokemonBossArena bằng Unity 6000.5.7f1, chờ nhập tài nguyên và biên dịch xong. Mở Assets/Scenes/MainMenu.unity rồi nhấn Play. Chọn một boss để bắt đầu; dùng mục Luyện tập khi cần thử riêng kỹ năng. Nếu Unity chặn Play Mode, kiểm tra lỗi đỏ trong Console trước khi xử lý cảnh báo thông thường.')
p('Các module hiện khai báo gồm audio, imgui, jsonserialize và screencapture; imageconversion được kéo theo phụ thuộc. Giữ nguyên metadata XML, texture Shadow và file meta của tài nguyên để tránh lỗi khung hình hoặc điểm neo. Scene chiến đấu phải được đăng ký trong Build Settings.')
h('Thêm boss mới')
p('Tạo lớp boss theo giao diện ArenaBoss, xây dựng bốn kỹ năng cùng vùng cảnh báo và hiệu ứng; chuẩn bị scene và giới hạn sân. Bổ sung cơ chế sinh đúng boss trong ArenaPreview, đăng ký scene, rồi thêm mục vào BossCatalog.json. Cuối cùng kiểm tra HP, hồi chiêu, sát thương, cách né, tương tác với chiêu điện, thắng thua, luyện tập và chơi lại. Thêm thẻ menu chỉ là bước đăng ký nội dung, không thay thế phần lập trình chiến đấu.')
h('Tài liệu tham chiếu nội bộ')
table(['Tài liệu hoặc mã nguồn','Nội dung sử dụng trong báo cáo'],[
('README.md và Docs/REQUIREMENTS.md','Mục tiêu, công nghệ, phạm vi và yêu cầu.'),
('Docs/PIKACHU_ABILITIES.md và PIKACHU_SOURCE.md','Điều khiển, kỹ năng, tài nguyên, sửa hoạt ảnh và nguồn sprite.'),
('Docs/CHARIZARD_BOSS.md và BLASTOISE_ARENA.md','Cơ chế boss, sân đấu và bằng chứng kiểm thử lịch sử.'),
('Docs/VENUSAUR_ARENA.md','Thông số màn ba, logic chiêu và xác nhận kiểm thử.'),
('Docs/SPELL_CLASHES.md và BOSS_SELECTION.md','Va chạm chiêu, quy tắc sức mạnh và menu mở rộng.'),
('Assets/Scripts và Assets/Resources/BossCatalog.json','Đối chiếu hành vi, thông số và danh sách boss hiện tại.'),
('Docs/Previews và Packages/manifest.json','Log, hình kiểm thử và phụ thuộc module Unity.')],[7.3,9.3])
p('Nguồn tài nguyên hình ảnh được ghi trong hồ sơ dự án: PMDCollab tại https://sprites.pmdcollab.org/ và kho SpriteCollab. Credit của từng Pokémon được lưu cùng thư mục Resources tương ứng. Hình minh họa trong báo cáo là ảnh của chính dự án.')
h('Đánh giá tổng thể')
p('Dự án đã có nền tảng chiến đấu và vòng chơi hoàn chỉnh ở mức prototype, với ba boss khác nhau và kiến trúc đủ để tiếp tục mở rộng. Bước tiếp theo quan trọng là hợp nhất dữ liệu cân bằng, kiểm thử hồi quy cho thay đổi mới, đo hiệu năng và hoàn thiện phản hồi nghe nhìn trước khi tăng nhanh số lượng boss.')

doc.save(OUT)
print(OUT)
